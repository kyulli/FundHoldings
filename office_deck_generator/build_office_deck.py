from __future__ import annotations

from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from office_deck_generator.config import (
    OFFICE_TEMPLATE,
    DEFAULT_DECK_OUTPUT,
    OUTPUT_DIR,
    PIPELINE_SUMMARY_METRICS,
)

from office_deck_generator.load_outputs import discover_outputs
from office_deck_generator.extract_tables import read_sheet_if_exists


# Office-style presentation colors
COLOR_DARK_BROWN = RGBColor(92, 51, 45)
COLOR_DEEP_RED = RGBColor(128, 38, 38)

COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_WARM_WHITE = RGBColor(250, 247, 244)
COLOR_LIGHT_TAN = RGBColor(239, 229, 220)

COLOR_TEXT = RGBColor(55, 45, 42)
COLOR_BORDER = RGBColor(190, 170, 155)


def format_percent(value) -> str:
    """Format a decimal ratio as a percentage for display only."""
    if pd.isna(value):
        return ""

    return f"{float(value):.1%}"


def friendly_label(value: str) -> str:
    """Convert machine-readable enum labels into presentation text."""
    if not isinstance(value, str):
        return str(value)

    return value.replace("_", " ").title()


def clear_existing_slides(prs: Presentation) -> None:
    """
    Remove example/template slides while preserving layouts and masters.
    """

    slide_ids = list(prs.slides._sldIdLst)

    for slide_id in slide_ids:
        prs.part.drop_rel(slide_id.rId)
        prs.slides._sldIdLst.remove(slide_id)


def set_placeholder_text(
    slide,
    idx: int,
    text: str,
) -> None:
    """
    Set text on a placeholder by placeholder index.
    """

    for placeholder in slide.placeholders:
        if placeholder.placeholder_format.idx == idx:
            placeholder.text = text
            return


def add_title_slide(
    prs: Presentation,
    title: str,
    subtitle: str,
):
    slide = prs.slides.add_slide(
        prs.slide_layouts[5]
    )

    set_placeholder_text(
        slide,
        0,
        f"{title}\n{subtitle}",
    )

    return slide


def add_section_slide(
    prs: Presentation,
    title: str,
):
    slide = prs.slides.add_slide(
        prs.slide_layouts[4]
    )

    set_placeholder_text(
        slide,
        1,
        title,
    )

    return slide


def add_image_slide(
    prs: Presentation,
    title: str,
    image_path: Path,
    note: str = "",
):
    slide = prs.slides.add_slide(
        prs.slide_layouts[0]
    )

    set_placeholder_text(
        slide,
        0,
        title,
    )

    content_placeholder = None

    for placeholder in slide.placeholders:
        if placeholder.placeholder_format.idx == 1:
            content_placeholder = placeholder
            break

    if content_placeholder is None:
        raise RuntimeError(
            "Content placeholder idx=1 not found."
        )

    left = content_placeholder.left
    top = content_placeholder.top
    width = content_placeholder.width
    height = content_placeholder.height

    # Remove the unused content placeholder shape.
    sp = content_placeholder._element
    sp.getparent().remove(sp)

    slide.shapes.add_picture(
        str(image_path),
        left,
        top,
        width=width,
        height=height,
    )

    if note:
        set_placeholder_text(
            slide,
            13,
            note,
        )

    return slide


def style_table(
    table,
) -> None:
    """
    Apply a red / white / brown financial-reporting style.
    """

    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.margin_left = Pt(7)
            cell.margin_right = Pt(7)
            cell.margin_top = Pt(5)
            cell.margin_bottom = Pt(5 )

            # Header row
            if row_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLOR_DARK_BROWN

                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = COLOR_WHITE
                        run.font.bold = True
                        run.font.size = Pt(11)

            # Alternating body rows
            else:
                cell.fill.solid()

                if row_idx % 2 == 0:
                    cell.fill.fore_color.rgb = COLOR_WARM_WHITE
                else:
                    cell.fill.fore_color.rgb = COLOR_LIGHT_TAN

                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = COLOR_TEXT
                        run.font.size = Pt(10.5)


def dataframe_to_slide_table(
    slide,
    df: pd.DataFrame,
    max_rows: int = 12,
    percent_columns: list[str] | None = None,
    friendly_columns: list[str] | None = None,
):
    """
    Insert existing DataFrame values into the template content area.
    """

    percent_columns = percent_columns or []
    friendly_columns = friendly_columns or []

    df = df.head(max_rows).copy()

    content_placeholder = None

    for placeholder in slide.placeholders:
        if placeholder.placeholder_format.idx == 1:
            content_placeholder = placeholder
            break

    if content_placeholder is None:
        raise RuntimeError(
            "Content placeholder idx=1 not found."
        )

    left = content_placeholder.left
    top = content_placeholder.top
    width = content_placeholder.width
    height = content_placeholder.height

    sp = content_placeholder._element
    sp.getparent().remove(sp)

    rows = len(df) + 1
    cols = len(df.columns)

    table_shape = slide.shapes.add_table(
        rows,
        cols,
        left,
        top,
        width,
        height,
    )

    table = table_shape.table

    for col_idx, column in enumerate(df.columns):
        table.cell(0, col_idx).text = str(column)

    for row_idx, (_, row) in enumerate(
        df.iterrows(),
        start=1,
    ):
        for col_idx, value in enumerate(row):
            column_name = df.columns[col_idx]

            if pd.isna(value):
                text = ""

            elif column_name in percent_columns:
                text = format_percent(value)

            elif column_name in friendly_columns:
                text = friendly_label(value)

            elif isinstance(value, float):
                text = f"{value:.3f}".rstrip("0").rstrip(".")

            else:
                text = str(value)

            table.cell(
                row_idx,
                col_idx,
            ).text = text

    # Formatting
    style_table(table)

    return table


def add_table_slide(
    prs: Presentation,
    title: str,
    df: pd.DataFrame,
    note: str = "",
    max_rows: int = 12,
    percent_columns: list[str] | None = None,
    friendly_columns: list[str] | None = None,
    highlight_severity: bool = False,
):
    slide = prs.slides.add_slide(
        prs.slide_layouts[0]
    )

    set_placeholder_text(
        slide,
        0,
        title,
    )

    table = dataframe_to_slide_table(
        slide,
        df,
        max_rows=max_rows,
        percent_columns=percent_columns,
        friendly_columns=friendly_columns,
    )

    if note:
        set_placeholder_text(
            slide,
            13,
            note,
        )

    return slide, table


def add_workflow_slide(
    prs: Presentation,
):
    """
    Add a readable Office-style end-to-end workflow slide.
    """

    slide = prs.slides.add_slide(
        prs.slide_layouts[0]
    )

    set_placeholder_text(
        slide,
        0,
        "End-to-End Data Quality Workflow",
    )

    content_placeholder = None

    for placeholder in slide.placeholders:
        if placeholder.placeholder_format.idx == 1:
            content_placeholder = placeholder
            break

    if content_placeholder is None:
        raise RuntimeError(
            "Content placeholder idx=1 not found."
        )

    left = content_placeholder.left
    top = content_placeholder.top
    width = content_placeholder.width
    height = content_placeholder.height

    # Remove placeholder
    sp = content_placeholder._element
    sp.getparent().remove(sp)


    stages = [
        {
            "number": "01",
            "title": "Data-State Analysis",
            "description": (
                "Assess completeness, consistency, "
                "and fund- and manager-level "
                "data quality."
            ),
        },

        {
            "number": "02",
            "title": "Quarterly Cleaning Pipeline",
            "description": (
                "Refresh validation outputs and "
                "compare results against prior "
                "reporting baselines."
            ),
        },

        {
            "number": "03",
            "title": "PDF Validation",
            "description": (
                "Compare structured holdings "
                "with source-report evidence "
                "from validated PDFs."
            ),
        },

        {
            "number": "04",
            "title": "Entity Resolution",
            "description": (
                "Confirm entity relationships "
                "between PDF disclosures and "
                "structured holdings."
            ),
        },
        
        {
            "number": "05",
            "title": "Exception Diagnosis",
            "description": (
                "Classify validation findings "
                "and provide Office review "
                "guidance."
            ),
        },
    ]


    # Layout: 3 cards top, 2 cards bottom
    card_width = int(width * 0.28)
    card_height = int(height * 0.34)

    horizontal_gap = int(width * 0.035)
    vertical_gap = int(height * 0.12)

    row1_top = (
        top
        + int(height * 0.08)
    )

    row2_top = (
        row1_top
        + card_height
        + vertical_gap
    )

    row1_total_width = (
        card_width * 3
        + horizontal_gap * 2
    )

    row1_left = (
        left
        + int(
            (width - row1_total_width)
            / 2
        )
    )

    positions = [

        # 01
        (
            row1_left,
            row1_top,
        ),

        # 02
        (
            row1_left
            + card_width
            + horizontal_gap,
            row1_top,
        ),

        # 03
        (
            row1_left
            + (card_width + horizontal_gap) * 2,
            row1_top,
        ),

        # 04
        (
            left
            + int(width * 0.22),
            row2_top,
        ),

        # 05
        (
            left
            + int(width * 0.55),
            row2_top,
        ),
    ]


    # Draw cards
    for i, stage in enumerate(stages):

        card_left, card_top = positions[i]

        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            card_left,
            card_top,
            card_width,
            card_height,
        )

        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_WHITE

        card.line.color.rgb = COLOR_BORDER
        card.line.width = Pt(1.25)


        # Number tag
        tag_width = int(
            card_width * 0.22
        )

        tag_height = int(
            card_height * 0.18
        )

        tag = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            card_left
            + int(card_width * 0.07),

            card_top
            + int(card_height * 0.08),

            tag_width,
            tag_height,
        )

        tag.fill.solid()
        tag.fill.fore_color.rgb = COLOR_DEEP_RED
        tag.line.fill.background()

        tag_tf = tag.text_frame
        tag_tf.clear()
        tag_tf.vertical_anchor = MSO_ANCHOR.MIDDLE

        p = tag_tf.paragraphs[0]
        p.text = stage["number"]
        p.alignment = PP_ALIGN.CENTER

        run = p.runs[0]
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = COLOR_WHITE


        # Title
        title_box = slide.shapes.add_textbox(
            card_left
            + int(card_width * 0.07),

            card_top
            + int(card_height * 0.32),
            int(card_width * 0.86),
            int(card_height * 0.25),
        )

        title_tf = title_box.text_frame
        title_tf.clear()
        title_tf.word_wrap = True

        p = title_tf.paragraphs[0]
        p.text = stage["title"]

        run = p.runs[0]
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = COLOR_DARK_BROWN


        # Description
        desc_box = slide.shapes.add_textbox(
            card_left
            + int(card_width * 0.07),

            card_top
            + int(card_height * 0.62),
            int(card_width * 0.86),
            int(card_height * 0.25),
        )

        desc_tf = desc_box.text_frame
        desc_tf.clear()
        desc_tf.word_wrap = True

        p = desc_tf.paragraphs[0]
        p.text = stage["description"]

        run = p.runs[0]
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_TEXT


    # Add arrows
    arrow_pairs = [
        (0, 1),
        (1, 2),
        (3, 4),
    ]

    for start_idx, end_idx in arrow_pairs:

        start_x = (
            positions[start_idx][0]
            + card_width
            - int(horizontal_gap * 0.25)
        )

        start_y = (
            positions[start_idx][1]
            + int(card_height * 0.45)
        )

        arrow_width = (
            positions[end_idx][0]
            - start_x
        )

        arrow = slide.shapes.add_textbox(
            start_x,
            start_y,
            arrow_width,
            int(card_height * 0.15),
        )

        p = arrow.text_frame.paragraphs[0]
        p.text = "→"
        p.alignment = PP_ALIGN.CENTER

        run = p.runs[0]
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = COLOR_DARK_BROWN


    set_placeholder_text(
        slide,
        13,
        (
            "Each stage builds on upstream outputs, "
            "transforming data validation results "
            "into actionable Office review guidance."
        ),
    )

    return slide


def build_deck():
    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    outputs = discover_outputs()

    prs = Presentation(
        OFFICE_TEMPLATE
    )

    clear_existing_slides(prs)


    # Title
    add_title_slide(
        prs,
        "Private Funds Data Quality & Validation",
        "August 2026",
    )


    # Workflow
    add_workflow_slide(
        prs
    )


        # Phase 1: Data-State Analysis

    phase1_summary = outputs.get(
        "phase1_summary",
    )


    if phase1_summary:

        # Slide 1: Executive Data-State Overview

        overview_rows = []

        coverage = phase1_summary.get(
            "coverage",
            {}
        )

        for key in [
            "Funds Covered",
            "Managers Covered",
        ]:
            if key in coverage:
                overview_rows.append(
                    {
                        "Metric": key,
                        "Value": coverage[key],
                    }
                )


        completeness = phase1_summary.get(
            "completeness",
            {}
        )

        for key in [
            "Unconditional Completeness",
            "Conditional Completeness",
            "Conditional Improvement",
        ]:
            if key in completeness:
                overview_rows.append(
                    {
                        "Metric": key,
                        "Value": completeness[key],
                    }
                )

        overview_df = pd.DataFrame(
            overview_rows
        )

        if not overview_df.empty:

            add_table_slide(
                prs,
                "Phase 1 Data-State Overview",
                overview_df,
                note=(
                    "Phase 1 assessed current investment "
                    "data coverage and completeness. "
                    "Conditional completeness adjusts for "
                    "fields not applicable to certain "
                    "investment scenarios."
                ),
                max_rows=10,
            )


        # Slide 2: Key Data Quality Observations

        observation_rows = []

        # Fund completeness

        fund_quality = phase1_summary.get(
            "fund_quality",
            {}
        )

        if fund_quality:

            unconditional = fund_quality.get(
                "Funds Below 85% Unconditional"
            )

            conditional = fund_quality.get(
                "Funds Below 85% Conditional"
            )

            if (
                unconditional is not None
                and conditional is not None
            ):

                observation_rows.append(
                    {
                        "Area":
                            "Fund Completeness",

                        "Finding":
                            (
                                f"Funds below 85% "
                                f"completeness decreased "
                                f"from {unconditional} "
                                f"to {conditional} "
                                f"after excluding "
                                f"non-applicable fields."
                            ),
                    }
                )


        # Missing fields

        field_quality = phase1_summary.get(
            "field_quality",
            []
        )

        if field_quality:

            top_fields = ", ".join(
                [
                    item["Field"]
                    for item in field_quality[:2]
                ]
            )

            observation_rows.append(
                {
                    "Area":
                        "Missing Fields",

                    "Finding":
                        (
                            f"{top_fields} "
                            "represent the largest "
                            "conditional reporting gaps."
                        ),
                }
            )


        # Manager quality

        manager_quality = phase1_summary.get(
            "manager_quality",
            {}
        )

        if manager_quality:

            observation_rows.append(
                {
                    "Area":
                        "Manager Quality",

                    "Finding":
                        (
                            f"{manager_quality.get('Managers Scored')} "
                            "managers evaluated with "
                            f"{manager_quality.get('Average Conditional Completeness')} "
                            "average conditional completeness."
                        ),
                }
            )


        # Data hygiene

        data_hygiene = phase1_summary.get(
            "data_hygiene",
            {}
        )

        if data_hygiene:

            observation_rows.append(
                {
                    "Area":
                        "Data Hygiene",

                    "Finding":
                        (
                            f"{data_hygiene.get('Excluded Non-Investment Rows')} "
                            "non-investment rows excluded "
                            "before quality assessment."
                        ),
                }
            )

        observations_df = pd.DataFrame(
            observation_rows
        )

        if not observations_df.empty:

            add_table_slide(
                prs,
                "Key Data Quality Observations",
                observations_df,
                note=(
                    "Summary of the primary reporting "
                    "gaps and data quality considerations "
                    "identified during Phase 1."
                ),
                max_rows=10,
            )

  
    data_state_images = outputs.get(
        "data_state_images",
        [],
    )

    if data_state_images:
        add_section_slide(
            prs,
            "I. Data-State Analysis",
        )

    image_titles = {
        "completeness_by_dimension.png":
            "Data Completeness by Dimension",

        "fund_quality_distribution_outliers.png":
            "Fund Quality Distribution & Outliers",

        "manager_field_heatmap.png":
            "Manager-Level Field Completeness",

        "portfolio_concentration.png":
            "Portfolio Concentration",

        "quality_by_investment_type.png":
            "Data Quality by Investment Type",
    }

    image_notes = {
        "completeness_by_dimension.png":
            (
                "Conditional completeness improves reporting coverage "
                "by excluding fields that are not applicable to certain "
                "investment scenarios."
            ),

        "fund_quality_distribution_outliers.png":
            (
                "Fund-level completeness varies across the portfolio. "
                "After excluding non-applicable fields, funds below the "
                "85% threshold decreased from 102 to 83."
            ),

        "manager_field_heatmap.png":
            (
                "Manager-level analysis highlights differences in "
                "reporting completeness across investment managers. "
                "89 managers were included in the conditional assessment."
            ),

        "portfolio_concentration.png":
            (
                "Portfolio concentration provides additional context "
                "when interpreting data quality patterns across "
                "investment exposures."
            ),

        "quality_by_investment_type.png":
            (
                "Data quality varies by investment type due to "
                "differences in reporting structures and available "
                "investment attributes."
            ),
    }

    for image_path in data_state_images:
        title = image_titles.get(
            image_path.name,
            image_path.stem.replace("_", " ").title(),
        )

        add_image_slide(
            prs,
            title,
            image_path,
            note=image_notes.get(
                image_path.name,
                ""
            ),
        )


    # Cleaning Pipeline
    ops_report = outputs.get(
        "ops_report"
    )

    if ops_report:
        add_section_slide(
            prs,
            "II. Quarterly Cleaning Pipeline",
        )

        summary_df = read_sheet_if_exists(
            ops_report,
            "Summary",
        )

        if summary_df is not None:
            summary_df = summary_df.dropna(
                how="all"
            )

            summary_df = summary_df[
                summary_df["Metric"].isin(
                    PIPELINE_SUMMARY_METRICS
                )
            ].copy()

            add_table_slide(
                prs,
                "Quarterly Pipeline Summary",
                summary_df,
                note=(
                    "Source: current Ops report. "
                    "Selected summary metrics are reproduced directly "
                    "from the existing workbook; full issue-level detail "
                    "remains available in the Ops report."
                ),
                max_rows=10,
            )


    # PDF Validation

    pdf_metrics = outputs.get(
        "pdf_validation_metrics",
    )

    if pdf_metrics:

        add_section_slide(
            prs,
            "III. PDF Validation",
        )


        # PDF Validation Overview

        pdf_metrics_df = pd.DataFrame(
            [
                {
                    "Validation Area": key,
                    "Result": value,
                }
                for key, value in pdf_metrics.items()
            ]
        )

        add_table_slide(
            prs,
            "PDF Validation Overview",
            pdf_metrics_df,
            note=(
                "PDF validation assesses whether source documents "
                "can be reliably compared against structured "
                "holdings data."
            ),
            max_rows=10,
        )


        # PDF Validation Insights

        pdf_insights_df = pd.DataFrame(
            [
                {
                    "Area":
                        "Document Comparability",

                    "Finding":
                        (
                            "All reviewed PDFs passed "
                            "comparability checks and were "
                            "available for automated validation."
                        ),
                },

                {
                    "Area":
                        "Value Validation",

                    "Finding":
                        (
                            "No PDF / vendor value differences "
                            "were identified in reviewed samples."
                        ),
                },

                {
                    "Area":
                        "Review Queue",

                    "Finding":
                        (
                            "Remaining documents require "
                            "targeted review based on "
                            "validation conditions."
                        ),
                },
            ]
        )

        add_table_slide(
            prs,
            "PDF Validation Insights",
            pdf_insights_df,
            note=(
                "Summary of document validation outcomes "
                "and remaining review considerations."
            ),
            max_rows=10,
        )


    # Entity Resolution

    entity_metrics = outputs.get(
        "entity_resolution_metrics",
    )

    if entity_metrics:

        add_section_slide(
            prs,
            "IV. Entity Resolution",
        )


        # Entity Resolution Overview

        entity_metrics_df = pd.DataFrame(
            [
                {
                    "Metric": key,
                    "Value": value,
                }
                for key, value in entity_metrics.items()
            ]
        )

        add_table_slide(
            prs,
            "Entity Resolution Overview",
            entity_metrics_df,
            note=(
                "Entity resolution links PDF-extracted company "
                "names with structured holdings records "
                "before downstream validation."
            ),
            max_rows=10,
        )


        # Entity Resolution Insights

        entity_insights_df = pd.DataFrame(
            [
                {
                    "Area":
                        "Entity Linking",

                    "Finding":
                        (
                            "PDF company names were matched "
                            "against structured holdings entities."
                        ),
                },

                {
                    "Area":
                        "Mapping Coverage",

                    "Finding":
                        (
                            "Confirmed mappings were generated "
                            "across reviewed source documents."
                        ),
                },

                {
                    "Area":
                        "Remaining Review",

                    "Finding":
                        (
                            "Unconfirmed mappings require "
                            "review before automated comparison."
                        ),
                },
            ]
        )

        add_table_slide(
            prs,
            "Entity Resolution Insights",
            entity_insights_df,
            note=(
                "Entity resolution provides the mapping layer "
                "required for reliable PDF-to-data validation."
            ),
            max_rows=10,
        )


    # Diagnosis
    diagnosis_tables = outputs.get(
        "diagnosis_tables",
        {}
    )

    if diagnosis_tables:

        add_section_slide(
            prs,
            "V. Exception Diagnosis",
        )

        # Exception Diagnosis Summary
        exception_summary = diagnosis_tables.get(
            "exception_summary"
        )

        if (
            exception_summary is not None
            and not exception_summary.empty
        ):

            add_table_slide(
                prs,
                "Exception Diagnosis Summary",
                exception_summary[
                    [
                        "Exception Type",
                        "Count",
                        "Diagnosis",
                        "Affected Funds",
                        "Affected Managers",
                    ]
                ],
                note=(
                    "Source: diagnosis report. "
                    "Exceptions are classified based on "
                    "validation results and available evidence. "
                    "Impact reflects affected funds and managers."
                ),
                max_rows=10,
            )


        # Recommendation
        recommendation_source = diagnosis_tables.get(
            "exception_summary"
        )

        if (
            recommendation_source is not None
            and not recommendation_source.empty
        ):

            recommendation_df = (
                recommendation_source[
                    [
                        "Exception Type",
                        "Recommended Action",
                        "Review Guidance",
                    ]
                ]
                .drop_duplicates()
            )

            add_table_slide(
                prs,
                "Recommended Review Actions",
                recommendation_df,
                note=(
                    "Recommended actions provide Office guidance "
                    "for reviewing and resolving identified exceptions."
                ),
                max_rows=10,
            )

    prs.save(
        DEFAULT_DECK_OUTPUT
    )

    print(
        f"Deck written to: {DEFAULT_DECK_OUTPUT}"
    )


if __name__ == "__main__":
    build_deck()