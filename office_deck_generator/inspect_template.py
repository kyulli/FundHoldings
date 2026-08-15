from pptx import Presentation

from office_deck_generator.config import OFFICE_TEMPLATE


def inspect_template():
    prs = Presentation(OFFICE_TEMPLATE)

    print(f"Template: {OFFICE_TEMPLATE}")
    print(f"Slide size: {prs.slide_width} x {prs.slide_height}")
    print(f"Number of existing slides: {len(prs.slides)}")
    print(f"Number of slide layouts: {len(prs.slide_layouts)}")

    print("\n=== SLIDE LAYOUTS ===")

    for i, layout in enumerate(prs.slide_layouts):
        print(f"\nLayout {i}: {layout.name}")

        for placeholder in layout.placeholders:
            ph_format = placeholder.placeholder_format

            print(
                "  "
                f"idx={ph_format.idx}, "
                f"type={ph_format.type}, "
                f"name={placeholder.name}"
            )

    print("\n=== EXISTING TEMPLATE SLIDES ===")

    for i, slide in enumerate(prs.slides, start=1):
        print(f"\nSlide {i}")

        for shape in slide.shapes:
            text = ""

            if hasattr(shape, "text"):
                text = shape.text.strip().replace("\n", " | ")

            print(
                f"  {shape.shape_type}: "
                f"{shape.name}"
                + (f" -> {text[:120]}" if text else "")
            )


if __name__ == "__main__":
    inspect_template()